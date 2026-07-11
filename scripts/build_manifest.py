#!/usr/bin/env python3
"""Build the deterministic accepted-paper run manifest."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _csv_rows(path: Path) -> int:
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.reader(handle)
        try:
            next(reader)
        except StopIteration:
            return 0
        return sum(1 for _ in reader)


def _dimensions(path: Path) -> dict[str, int]:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8", newline="") as handle:
            reader = csv.reader(handle)
            try:
                header = next(reader)
            except StopIteration:
                return {"rows": 0, "columns": 0}
        return {"rows": _csv_rows(path), "columns": len(header)}
    if path.suffix.lower() == ".ipynb":
        notebook = json.loads(path.read_text(encoding="utf-8"))
        return {"cells": len(notebook.get("cells", []))}
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if path.name == "NNT_LRs_08-26-2025.xlsx":
                worksheet = workbook["Master"]
                return {
                    "sheets": len(workbook.sheetnames),
                    "rows": worksheet.max_row - 1,
                    "columns": worksheet.max_column,
                }
            rows = sum(max(worksheet.max_row - 2, 0) for worksheet in workbook.worksheets)
            columns = max(worksheet.max_column for worksheet in workbook.worksheets)
            return {"sheets": len(workbook.sheetnames), "rows": rows, "columns": columns}
        finally:
            workbook.close()
    return {}


def _artifact(root: Path, relative_path: str, role: str) -> dict[str, Any]:
    path = root / relative_path
    if not path.is_file():
        raise FileNotFoundError(f"Required manifest artifact is missing: {relative_path}")
    artifact: dict[str, Any] = {
        "path": relative_path,
        "sha256": sha256_file(path),
        "byte_size": path.stat().st_size,
        "role": role,
    }
    dimensions = _dimensions(path)
    if dimensions:
        artifact["dimensions"] = dimensions
    return artifact


def build_manifest(root: Path, output_path: Path) -> Path:
    """Build and write the release manifest."""
    root = root.resolve()
    output_path = output_path.resolve()

    configuration_specs = (
        (".python-version", "python_version_pin"),
        ("pyproject.toml", "primary_python_environment"),
        ("config/analysis_categories_v1.json", "qualitative_lr_category_definition"),
        ("config/manuscript_models_v1.json", "accepted_manuscript_model_scope"),
        ("manifests/manifest.schema.json", "manifest_validation_schema"),
        ("tools/cff/.python-version", "cff_validator_python_pin"),
        ("tools/cff/pyproject.toml", "cff_validator_environment"),
        ("tools/cff/uv.lock", "cff_validator_lockfile"),
    )
    prompt_specs = (
        ("prompts/main_estimator_v1.json", "accepted_main_estimator_prompt"),
        ("prompts/threshold_perturbation_v1.json", "accepted_threshold_perturbation_prompt"),
    )
    input_specs = (
        ("NNT_LRs_08-26-2025.xlsx", "accepted_analysis_workbook"),
        ("nnt_lrs_with_estimated.xlsx", "per_condition_labels_and_historical_outputs"),
        ("data/curated/diagnostic_lrs_manuscript_v1.csv", "canonical_accepted_analysis_data"),
        ("data/model_outputs/manuscript_model_outputs_v1.csv", "accepted_three_model_outputs"),
        ("data/model_outputs/auxiliary_gpt-4.1_outputs_v1.csv", "auxiliary_historical_not_analyzed"),
        ("data/provenance/source_crosswalk_v1.csv", "strict_workbook_crosswalk"),
    )
    threshold_base = "data/model_outputs/threshold_perturbation_v1"
    output_specs = (
        ("data/model_outputs/manuscript_query_run_v1.csv", "recoverable_main_query_configuration"),
        ("data/provenance/curation_log_v1.csv", "recoverable_row_level_curation_log"),
        ("data/provenance/provenance_gaps_v1.csv", "known_provenance_gaps"),
        (f"{threshold_base}/threshold_perturbation_cases.csv", "threshold_perturbation_cases"),
        (f"{threshold_base}/threshold_perturbation_raw.csv", "accepted_threshold_raw_outputs"),
        (f"{threshold_base}/threshold_perturbation_summary.csv", "threshold_perturbation_summary"),
        (
            f"{threshold_base}/threshold_perturbation_reviewer_table.csv",
            "threshold_perturbation_response_table",
        ),
        (f"{threshold_base}/threshold_perturbation_summary.md", "threshold_perturbation_narrative_summary"),
        ("results/reference/main_metrics.json", "dataset_reference_contract"),
        ("results/reference/category_counts.csv", "reported_lr_category_counts"),
        ("results/reference/feature_type_counts.csv", "overlapping_feature_memberships"),
        ("results/reference/agreement_metrics.csv", "bland_altman_and_log_error_metrics"),
        ("results/reference/pairwise_model_comparisons.csv", "paired_bias_and_dispersion_comparisons"),
        ("results/reference/coverage_intervals.csv", "multiplicative_coverage_intervals"),
        ("results/reference/evidence_direction_tests.csv", "negative_vs_positive_evidence_tests"),
        ("results/reference/calibration_metrics.csv", "primary_calibration_metrics"),
        ("results/reference/reliability_metrics.csv", "inverse_reliability_calibration_metrics"),
        ("results/reference/kappa_metrics.csv", "qualitative_agreement_metrics"),
        ("results/reference/reliability_zone_metrics.csv", "central_and_extreme_zone_metrics"),
        (
            "results/reference/laboratory_discrepancy_flags.json",
            "gpt5_laboratory_discrepancy_flag_counts",
        ),
        ("llms-full.txt", "licensed_preproduction_author_manuscript_text"),
    )
    code_specs = (
        ("data_analysis.ipynb", "accepted_primary_analysis_notebook"),
        ("supplementary_analyses.ipynb", "accepted_supplementary_analysis_notebook"),
        ("lr_scraper_estimator.ipynb", "guarded_historical_collection_and_query_notebook"),
        (
            "threshold_perturbation_sensitivity_analysis.ipynb",
            "guarded_threshold_perturbation_notebook",
        ),
        ("scripts/build_release_dataset.py", "canonical_data_builder"),
        ("scripts/compute_reference_results.py", "reference_result_builder"),
        ("scripts/reproduce_paper.py", "offline_reproduction_entrypoint"),
        ("scripts/build_manifest.py", "manifest_builder"),
        ("scripts/build_checksums.py", "repository_checksum_builder"),
        ("scripts/verify_checksums.py", "repository_checksum_verifier"),
        ("scripts/validate_release.py", "release_contract_validator"),
        ("scripts/check_release_hygiene.py", "public_release_hygiene_gate"),
        ("scripts/run_replication.py", "guarded_live_replication_entrypoint"),
        ("scripts/extract_manuscript_text.py", "author_manuscript_text_extractor"),
        ("scripts/build_release_archive.py", "deterministic_source_archive_builder"),
        ("scripts/build_release_assets.py", "github_release_asset_builder"),
        ("scripts/generate_release_attestation.py", "release_attestation_builder"),
    )

    gaps_path = root / "data/provenance/provenance_gaps_v1.csv"
    with gaps_path.open("r", encoding="utf-8", newline="") as handle:
        gap_rows = list(csv.DictReader(handle))
    known_gaps = [
        {
            "gap_id": row["gap_id"],
            "artifact_or_field": row["artifact_or_field"],
            "status": row["status"],
            "blocking_for_reproduction": row["blocking_for_reproduction"].lower() == "true",
            "description": row["description"],
        }
        for row in gap_rows
    ]

    lockfile = root / "uv.lock"
    if not lockfile.is_file():
        raise FileNotFoundError("Required lockfile is missing: uv.lock")

    manifest = {
        "schema_version": 1,
        "manifest_id": "llm-estimate-lrs-manuscript-v1",
        "release_version": "1.0.0",
        "release_ref": "v1.0.0",
        "repository": "https://github.com/reblocke/llm_estimate_lrs",
        "article": {
            "title": (
                "Large language models generate diagnostic likelihood ratios with low mean bias but wide dispersion"
            ),
            "journal": "Scientific Reports",
            "doi": "10.1038/s41598-026-61766-2",
            "accepted_date": "2026-07-07",
            "published_date": "2026-07-11",
            "publication_status": "published_unedited_early_access",
            "publisher_url": "https://doi.org/10.1038/s41598-026-61766-2",
        },
        "dates": {"source_collection": "2025-04-01", "model_query": "2025-08-25"},
        "models_config": "config/manuscript_models_v1.json",
        "category_config": "config/analysis_categories_v1.json",
        "configurations": [_artifact(root, path, role) for path, role in configuration_specs],
        "prompts": [_artifact(root, path, role) for path, role in prompt_specs],
        "inputs": [_artifact(root, path, role) for path, role in input_specs],
        "outputs": [_artifact(root, path, role) for path, role in output_specs],
        "code": [_artifact(root, path, role) for path, role in code_specs],
        "environment": {
            "python_requirement": ">=3.11",
            "lockfile": "uv.lock",
            "lockfile_sha256": sha256_file(lockfile),
        },
        "reproduction_command": "make reproduce",
        "validation_command": "make release-check",
        "known_provenance_gaps": known_gaps,
        "release_attestation": "Generated separately from the approved tag and release archive.",
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(manifest, indent=2, sort_keys=True, ensure_ascii=False, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    return output_path


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path("."))
    parser.add_argument("--output", type=Path, default=Path("manifests/manuscript_run_v1.json"))
    return parser.parse_args()


def main() -> None:
    args = _parse_args()
    print(build_manifest(args.root, args.output))


if __name__ == "__main__":
    main()
