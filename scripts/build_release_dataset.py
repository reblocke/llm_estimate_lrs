#!/usr/bin/env python3
"""Build the frozen, machine-readable accepted-paper data release.

The builder never writes either source workbook. Crosswalking is exact after
minimal comparison-only normalization and uses an occurrence index to preserve
duplicate source rows.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import shutil
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

MASTER_SHA256 = "644f0558328a8f04f460a5ebfa2fc04e6d3571f655d084ea076488c7ba17da89"
PER_CONDITION_SHA256 = "c375229a27f0854957f6b8963ece145e94d00e4fa75e7e2756e5d130f6f7110d"

MANUSCRIPT_MODEL_COLUMNS = (
    "lr_gpt-4o-2024-11-20",
    "lr_o3-2025-04-16",
    "lr_gpt-5",
)
AUXILIARY_MODEL_COLUMN = "lr_gpt-4.1-2025-04-14"

THRESHOLD_FILES = {
    "threshold_perturbation_cases.csv": "0ecd978e3308f39a154ba709de28ac0bb5e7440f6160817d95de5d060d388c5b",
    "threshold_perturbation_raw.csv": "0f83fdc22140f75e86404a943ad90b039bd139eeb103db468522d49f825f70f6",
    "threshold_perturbation_reviewer_table.csv": "fda7aa0217587609e72f507f82a39f6657f9a02ac930d8c56b77af57ec96e7ba",
    "threshold_perturbation_summary.csv": "f0ab68b0ff51292aa0cc4a6f6191087ba3ed6827f06a5489ddd554b4c0ac6128",
    "threshold_perturbation_summary.md": "99fa0296700e30d5dbad3fe38e7b2d02f708dd5bf08c5fb49905c60a2b455c36",
}

EXPECTED_COLUMNS = (
    "finding",
    "lr_raw",
    "lr_reported",
    "lr_gpt-5",
    AUXILIARY_MODEL_COLUMN,
    "lr_gpt-4o-2024-11-20",
    "lr_o3-2025-04-16",
    "Feature Type",
    "condition",
)

CURATION_COLUMNS = (
    "curation_note_id",
    "row_id",
    "field",
    "raw_value",
    "curated_value",
    "reason",
    "reviewer_1",
    "reviewer_2",
    "source_reference",
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _require_hash(path: Path, expected: str) -> None:
    actual = sha256_file(path)
    if actual != expected:
        raise ValueError(f"Unexpected SHA-256 for {path}: expected {expected}, got {actual}")


def _normalise_text(value: Any) -> str:
    """Normalize only whitespace for comparison; exported raw text is untouched."""
    return re.sub(r"\s+", " ", str(value).strip())


def _normalise_number(value: Any) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Crosswalk numeric value is not numeric: {value!r}") from exc
    if not math.isfinite(number):
        raise ValueError(f"Crosswalk numeric value is not finite: {value!r}")
    return format(number, ".17g")


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _match_key(row: dict[str, Any], condition_field: str) -> tuple[str, ...]:
    return (
        _normalise_text(row[condition_field]),
        _normalise_text(row["finding"]),
        _normalise_number(row["lr_reported"]),
        _normalise_number(row["lr_gpt-4o-2024-11-20"]),
        _normalise_number(row["lr_o3-2025-04-16"]),
        _normalise_number(row["lr_gpt-5"]),
    )


def _read_master(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Master"]
        values = worksheet.iter_rows(values_only=True)
        header = tuple(next(values))
        if header != EXPECTED_COLUMNS:
            raise ValueError(f"Unexpected Master header: {header!r}")
        rows = []
        for master_row_number, values_row in enumerate(values, start=2):
            row = dict(zip(header, values_row, strict=True))
            row["master_row_number"] = master_row_number
            rows.append(row)
    finally:
        workbook.close()
    if len(rows) != 700:
        raise ValueError(f"Master sheet must contain 700 rows; found {len(rows)}")
    return rows


def _read_per_condition(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    try:
        if len(workbook.sheetnames) != 30:
            raise ValueError(f"Per-condition workbook must contain 30 sheets; found {len(workbook.sheetnames)}")
        for sheet_index, sheet_name in enumerate(workbook.sheetnames, start=1):
            worksheet = workbook[sheet_name]
            diagnosis = worksheet.cell(row=1, column=1).value
            if not isinstance(diagnosis, str) or not diagnosis.strip():
                raise ValueError(f"Missing full condition label in sheet {sheet_name!r}")
            header = tuple(cell.value for cell in worksheet[2])
            expected = EXPECTED_COLUMNS[:7]
            if header != expected:
                raise ValueError(f"Unexpected header in sheet {sheet_name!r}: {header!r}")
            for source_sheet_row, values_row in enumerate(worksheet.iter_rows(min_row=3, values_only=True), start=3):
                row = dict(zip(header, values_row, strict=True))
                row.update(
                    {
                        "source_sheet": sheet_name,
                        "source_sheet_row": source_sheet_row,
                        "source_sheet_index": sheet_index,
                        "condition_full": diagnosis.strip(),
                    }
                )
                rows.append(row)
    finally:
        workbook.close()
    if len(rows) != 700:
        raise ValueError(f"Per-condition workbook must contain 700 data rows; found {len(rows)}")
    return rows


def _classify_lr(value: float, categories: list[dict[str, Any]]) -> str:
    for category in categories:
        lower = category["lower"]
        upper = category["upper"]
        lower_ok = value > lower or (category["lower_inclusive"] and value == lower)
        upper_ok = upper is None or value < upper or (category["upper_inclusive"] and value == upper)
        if lower_ok and upper_ok:
            return str(category["label"])
    raise ValueError(f"Positive LR did not match a configured category: {value!r}")


def _source_lr_label(finding: str) -> str:
    normalized = _normalise_text(finding)
    normalized = normalized.replace("Patient does has:", "Patient does not have:", 1)
    if normalized.startswith("Patient does not have:"):
        return "LR−"
    if normalized.startswith("Patient has:"):
        return "LR+"
    raise ValueError(f"Cannot classify source LR direction from finding: {finding!r}")


def _write_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(path, index=False, lineterminator="\n", na_rep="")


def _build_crosswalk(
    master_rows: list[dict[str, Any]], per_rows: list[dict[str, Any]]
) -> tuple[list[tuple[dict[str, Any], dict[str, Any], int]], pd.DataFrame]:
    master_counts = Counter(_match_key(row, "condition") for row in master_rows)
    source_counts = Counter(_match_key(row, "source_sheet") for row in per_rows)
    if master_counts != source_counts:
        missing_source = list((master_counts - source_counts).items())[:10]
        extra_source = list((source_counts - master_counts).items())[:10]
        raise ValueError(
            "Master/per-condition crosswalk key multisets differ. "
            f"Missing source examples: {missing_source}; extra source examples: {extra_source}"
        )

    source_occurrences: defaultdict[tuple[str, ...], int] = defaultdict(int)
    source_lookup: dict[tuple[tuple[str, ...], int], dict[str, Any]] = {}
    for row in per_rows:
        key = _match_key(row, "source_sheet")
        source_occurrences[key] += 1
        source_lookup[(key, source_occurrences[key])] = row

    master_occurrences: defaultdict[tuple[str, ...], int] = defaultdict(int)
    matched: list[tuple[dict[str, Any], dict[str, Any], int]] = []
    crosswalk_rows = []
    for index, master_row in enumerate(master_rows, start=1):
        key = _match_key(master_row, "condition")
        master_occurrences[key] += 1
        occurrence = master_occurrences[key]
        source_row = source_lookup.get((key, occurrence))
        if source_row is None:
            raise ValueError(
                f"No per-condition match for Master row {master_row['master_row_number']} "
                f"at duplicate occurrence {occurrence}"
            )
        row_id = f"lr_{index:04d}"
        matched.append((master_row, source_row, occurrence))
        key_hash = hashlib.sha256(
            json.dumps(key, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        ).hexdigest()
        crosswalk_rows.append(
            {
                "row_id": row_id,
                "master_row_number": master_row["master_row_number"],
                "source_sheet": source_row["source_sheet"],
                "source_sheet_row": source_row["source_sheet_row"],
                "condition_full": source_row["condition_full"],
                "match_key_sha256": key_hash,
                "match_occurrence": occurrence,
                "master_key_count": master_counts[key],
                "source_key_count": source_counts[key],
                "match_status": "exact_occurrence_match",
            }
        )

    used = {(key, occurrence) for key, count in master_occurrences.items() for occurrence in range(1, count + 1)}
    if set(source_lookup) != used:
        raise ValueError("Crosswalk did not consume every per-condition row exactly once")
    return matched, pd.DataFrame(crosswalk_rows)


def _copy_threshold_artifacts(source_dir: Path, destination_dir: Path) -> dict[str, Path]:
    destination_dir.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    for filename, expected_hash in THRESHOLD_FILES.items():
        source = source_dir / filename
        if not source.is_file():
            raise FileNotFoundError(
                f"Blocking accepted-run artifact is missing: {source}. Do not recreate it with a live call."
            )
        _require_hash(source, expected_hash)
        destination = destination_dir / filename
        if source.resolve() != destination.resolve():
            shutil.copyfile(source, destination)
        _require_hash(destination, expected_hash)
        paths[filename] = destination
    return paths


def build_release_dataset(
    master_workbook: Path,
    per_condition_workbook: Path,
    output_root: Path,
    category_config: Path,
    threshold_source_dir: Path,
) -> dict[str, Path]:
    """Create deterministic release artifacts and return their paths."""
    master_workbook = master_workbook.resolve()
    per_condition_workbook = per_condition_workbook.resolve()
    output_root = output_root.resolve()
    category_config = category_config.resolve()
    threshold_source_dir = threshold_source_dir.resolve()

    _require_hash(master_workbook, MASTER_SHA256)
    _require_hash(per_condition_workbook, PER_CONDITION_SHA256)
    category_document = json.loads(category_config.read_text(encoding="utf-8"))
    categories = category_document["categories_in_order"]

    master_rows = _read_master(master_workbook)
    per_rows = _read_per_condition(per_condition_workbook)
    matched, crosswalk = _build_crosswalk(master_rows, per_rows)

    canonical_rows = []
    model_rows = []
    auxiliary_rows = []
    for index, (master, source, _occurrence) in enumerate(matched, start=1):
        row_id = f"lr_{index:04d}"
        numeric_values = {column: float(master[column]) for column in ("lr_reported", *MANUSCRIPT_MODEL_COLUMNS)}
        if any(not math.isfinite(value) or value <= 0 for value in numeric_values.values()):
            raise ValueError(f"Nonpositive or nonfinite accepted LR at {row_id}")
        feature_type = str(master["Feature Type"])
        reported = numeric_values["lr_reported"]
        canonical_rows.append(
            {
                "row_id": row_id,
                "master_row_number": master["master_row_number"],
                "source_sheet": source["source_sheet"],
                "source_sheet_row": source["source_sheet_row"],
                "condition_raw": _cell_text(master["condition"]),
                "condition_full": source["condition_full"],
                "finding_raw": _cell_text(master["finding"]),
                "finding_curated": _cell_text(master["finding"]),
                "lr_raw": _cell_text(master["lr_raw"]),
                "lr_reported": reported,
                "lr_gpt-4o-2024-11-20": numeric_values["lr_gpt-4o-2024-11-20"],
                "lr_o3-2025-04-16": numeric_values["lr_o3-2025-04-16"],
                "lr_gpt-5": numeric_values["lr_gpt-5"],
                "feature_type_raw": feature_type,
                "is_sign_or_symptom": bool(re.search(r"\b(sign|symptom)\b", feature_type, re.IGNORECASE)),
                "is_history": "history" in feature_type.lower(),
                "is_test_result": "test" in feature_type.lower(),
                "is_imaging": "imaging" in feature_type.lower(),
                "is_diagnostic_adjudication": "diagnosis" in feature_type.lower(),
                "source_lr_label": _source_lr_label(str(master["finding"])),
                "reported_stratum": "LR < 1" if reported < 1 else ("LR > 1" if reported > 1 else "LR = 1"),
                "reported_qualitative_band": _classify_lr(reported, categories),
                "source_url": "",
                "source_retrieved_date": "",
                "curation_status": "accepted_workbook_value",
                "curation_note_id": "",
            }
        )
        model_rows.append({"row_id": row_id, **{column: numeric_values[column] for column in MANUSCRIPT_MODEL_COLUMNS}})
        auxiliary_rows.append(
            {
                "row_id": row_id,
                AUXILIARY_MODEL_COLUMN: float(master[AUXILIARY_MODEL_COLUMN]),
                "accepted_manuscript_role": "auxiliary_historical_not_analyzed",
            }
        )

    canonical = pd.DataFrame(canonical_rows)
    if canonical["condition_full"].nunique() != 30:
        raise ValueError("Canonical data must contain exactly 30 full condition labels")

    output_paths: dict[str, Path] = {}
    output_paths["canonical"] = output_root / "data/curated/diagnostic_lrs_manuscript_v1.csv"
    output_paths["model_outputs"] = output_root / "data/model_outputs/manuscript_model_outputs_v1.csv"
    output_paths["auxiliary_outputs"] = output_root / "data/model_outputs/auxiliary_gpt-4.1_outputs_v1.csv"
    output_paths["query_audit"] = output_root / "data/model_outputs/manuscript_query_run_v1.csv"
    output_paths["crosswalk"] = output_root / "data/provenance/source_crosswalk_v1.csv"
    output_paths["curation_log"] = output_root / "data/provenance/curation_log_v1.csv"
    output_paths["provenance_gaps"] = output_root / "data/provenance/provenance_gaps_v1.csv"

    _write_csv(canonical, output_paths["canonical"])
    _write_csv(pd.DataFrame(model_rows), output_paths["model_outputs"])
    _write_csv(pd.DataFrame(auxiliary_rows), output_paths["auxiliary_outputs"])
    _write_csv(crosswalk, output_paths["crosswalk"])
    _write_csv(pd.DataFrame(columns=CURATION_COLUMNS), output_paths["curation_log"])

    query_audit = pd.DataFrame(
        [
            {
                "run_id": "",
                "query_date": "2025-08-25",
                "display_name": "GPT-4o",
                "api_model": "gpt-4o-2024-11-20",
                "prompt_id": "main-estimator-v1",
                "few_shot_set": "rich_8",
                "temperature": 0.2,
                "reasoning_effort": "",
                "text_verbosity": "",
                "max_retries": 8,
                "sdk_version": "",
                "response_ids_available": False,
                "retry_audit_available": False,
                "notes": "Original run identifier, response IDs, retry audit, and SDK version were not recovered.",
            },
            {
                "run_id": "",
                "query_date": "2025-08-25",
                "display_name": "o3",
                "api_model": "o3-2025-04-16",
                "prompt_id": "main-estimator-v1",
                "few_shot_set": "minimal_2",
                "temperature": "",
                "reasoning_effort": "medium",
                "text_verbosity": "",
                "max_retries": 8,
                "sdk_version": "",
                "response_ids_available": False,
                "retry_audit_available": False,
                "notes": "Original run identifier, response IDs, retry audit, and SDK version were not recovered.",
            },
            {
                "run_id": "",
                "query_date": "2025-08-25",
                "display_name": "GPT-5",
                "api_model": "gpt-5",
                "prompt_id": "main-estimator-v1",
                "few_shot_set": "minimal_2",
                "temperature": "",
                "reasoning_effort": "medium",
                "text_verbosity": "low",
                "max_retries": 8,
                "sdk_version": "",
                "response_ids_available": False,
                "retry_audit_available": False,
                "notes": (
                    "Alias used in the accepted run; original run identifier, response IDs, retry audit, "
                    "and SDK version were not recovered."
                ),
            },
        ]
    )
    _write_csv(query_audit, output_paths["query_audit"])

    provenance_gaps = pd.DataFrame(
        [
            {
                "gap_id": "PG-001",
                "artifact_or_field": "source_url",
                "scope": "row_level",
                "status": "not_recovered",
                "blocking_for_reproduction": False,
                "description": "Row-level TheNNT source URLs were not stored in the frozen workbooks.",
            },
            {
                "gap_id": "PG-002",
                "artifact_or_field": "source_retrieved_date",
                "scope": "row_level",
                "status": "not_recovered",
                "blocking_for_reproduction": False,
                "description": (
                    "The study-level source collection date is 2025-04-01; row-level retrieval "
                    "timestamps were not recorded."
                ),
            },
            {
                "gap_id": "PG-003",
                "artifact_or_field": "curation_log",
                "scope": "row_level",
                "status": "not_recovered",
                "blocking_for_reproduction": False,
                "description": "Row-level curation rationales and reviewer identities were not recovered.",
            },
            {
                "gap_id": "PG-004",
                "artifact_or_field": "response_ids",
                "scope": "main_model_queries",
                "status": "not_recovered",
                "blocking_for_reproduction": False,
                "description": "Original OpenAI response IDs were not recovered.",
            },
            {
                "gap_id": "PG-005",
                "artifact_or_field": "retry_audit",
                "scope": "main_model_queries",
                "status": "not_recovered",
                "blocking_for_reproduction": False,
                "description": (
                    "No original row-level retry audit was found in the worktree, local ignored files, or Git history."
                ),
            },
            {
                "gap_id": "PG-006",
                "artifact_or_field": "sdk_version",
                "scope": "main_model_queries",
                "status": "not_recovered",
                "blocking_for_reproduction": False,
                "description": "The exact OpenAI SDK version used for the accepted run was not recorded.",
            },
            {
                "gap_id": "PG-007",
                "artifact_or_field": "raw_api_responses",
                "scope": "main_model_queries",
                "status": "not_recovered",
                "blocking_for_reproduction": False,
                "description": (
                    "Raw API response JSON was not recovered; the accepted numeric outputs are frozen in the workbooks."
                ),
            },
            {
                "gap_id": "PG-008",
                "artifact_or_field": "query_timestamps",
                "scope": "main_model_queries",
                "status": "partially_recovered",
                "blocking_for_reproduction": False,
                "description": (
                    "The study-level query date is 2025-08-25; row-level query timestamps were not recorded."
                ),
            },
        ]
    )
    _write_csv(provenance_gaps, output_paths["provenance_gaps"])

    threshold_destination = output_root / "data/model_outputs/threshold_perturbation_v1"
    output_paths.update(
        {
            f"threshold_{name}": path
            for name, path in _copy_threshold_artifacts(threshold_source_dir, threshold_destination).items()
        }
    )
    return output_paths


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--master", type=Path, default=Path("NNT_LRs_08-26-2025.xlsx"))
    parser.add_argument("--per-condition", type=Path, default=Path("nnt_lrs_with_estimated.xlsx"))
    parser.add_argument("--output-root", type=Path, default=Path("."))
    parser.add_argument("--category-config", type=Path, default=Path("config/analysis_categories_v1.json"))
    parser.add_argument(
        "--threshold-source",
        type=Path,
        default=Path("results/2026-05-16/threshold_perturbation"),
    )
    return parser.parse_args()


def main() -> None:
    args = _parse_args()
    paths = build_release_dataset(
        master_workbook=args.master,
        per_condition_workbook=args.per_condition,
        output_root=args.output_root,
        category_config=args.category_config,
        threshold_source_dir=args.threshold_source,
    )
    for name, path in sorted(paths.items()):
        print(f"{name}: {path}")


if __name__ == "__main__":
    main()
